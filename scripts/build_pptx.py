"""Genera un PPTX con las diapositivas de INNBULTZADA y el video embebido en la slide 6.
Uso: python3 scripts/build_pptx.py"""
import os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

BRAND = 'public/brand'
OUT = 'entrega/INNBULTZADA-presentacion.pptx'
TMP = '/tmp/pptx_slides'
os.makedirs(TMP, exist_ok=True)
os.makedirs('entrega', exist_ok=True)

# Orden del deck y qué slides llevan el logo de LABORAL Kutxa (las de contenido).
ORDER = ['slide-1', 'slide-2', 'slide-3', 'slide-4', 'slide-5', 'slide-6',
         'slide-7', 'slide-7.1', 'slide-8', 'slide-9', 'slide-10']
CONTENT = {'slide-2', 'slide-3', 'slide-4', 'slide-5', 'slide-6', 'slide-7', 'slide-7.1', 'slide-8', 'slide-9'}
VIDEO_SLIDE = 'slide-6'

logo = Image.open(f'{BRAND}/LK_logo.png').convert('RGBA')

def composite(name):
    """Devuelve la ruta de la slide con el logo incrustado (si corresponde)."""
    im = Image.open(f'{BRAND}/{name}.png').convert('RGB')
    if name in CONTENT:
        W, H = im.size
        lw = int(W * 0.11); lh = int(logo.height * lw / logo.width)
        lg = logo.resize((lw, lh))
        im.paste(lg, (int(W * 0.03), int(H * 0.035)), lg)
    path = f'{TMP}/{name}.png'
    im.save(path)
    return path, im.size

# Poster del video (navy + triángulo play), en la proporción del video.
vw, vh = Image.open(f'{BRAND}/Duvan_aceleradoras.mp4') if False else (478, 850)
poster = Image.new('RGB', (478, 850), (18, 28, 51))
d = ImageDraw.Draw(poster)
cx, cy, r = 239, 425, 70
d.polygon([(cx - r + 20, cy - r), (cx - r + 20, cy + r), (cx + r + 10, cy)], fill=(255, 255, 255))
POSTER = f'{TMP}/poster.png'
poster.save(POSTER)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5
blank = prs.slide_layouts[6]

def contain(a):
    """Ubicación (x,y,w,h) en pulgadas para una imagen de aspecto a, centrada en 16:9 (fondo blanco)."""
    if a > SW / SH:
        w = SW; h = SW / a; return 0, (SH - h) / 2, w, h
    h = SH; w = SH * a; return (SW - w) / 2, 0, w, h

for name in ORDER:
    slide = prs.slides.add_slide(blank)
    # fondo blanco
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    path, (iw, ih) = composite(name)
    if name == VIDEO_SLIDE:
        # Video vertical embebido a la derecha; la imagen de la slide a la izquierda.
        vaspect = 478 / 850
        v_h = 6.0; v_w = v_h * vaspect
        v_x = SW - v_w - 0.35; v_y = (SH - v_h) / 2
        avail_w = v_x - 0.35 - 0.25
        i_w = avail_w; i_h = i_w * ih / iw
        if i_h > SH - 0.6:
            i_h = SH - 0.6; i_w = i_h * iw / ih
        i_x = 0.35; i_y = (SH - i_h) / 2
        slide.shapes.add_picture(path, Inches(i_x), Inches(i_y), Inches(i_w), Inches(i_h))
        slide.shapes.add_movie(f'{BRAND}/Duvan_aceleradoras.mp4', Inches(v_x), Inches(v_y),
                               Inches(v_w), Inches(v_h), poster_frame_image=POSTER, mime_type='video/mp4')
    else:
        x, y, w, h = contain(iw / ih)
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))

prs.save(OUT)
print(f'PPTX generado: {len(ORDER)} diapositivas → {OUT}')
